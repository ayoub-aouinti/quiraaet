#!/usr/bin/env python3
"""إخراج الملف المرقمن مفهرساً بالسور والآيات، بصيغة إكسل جاهزة للعمل."""
import json, sqlite3, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEAD = PatternFill('solid', fgColor='14584A')
NOTE = PatternFill('solid', fgColor='FFF4D6')
AR = 'Arial'


def sheet_head(ws, cols, widths):
    ws.sheet_view.rightToLeft = True
    ws.append(cols)
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(name=AR, bold=True, color='FFFFFF', size=11)
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'


def main(db, outfile):
    con = sqlite3.connect(db); con.row_factory = sqlite3.Row
    wb = Workbook()

    # --- ورقة الكلمات مع الفهرس ---
    ws = wb.active; ws.title = 'الكلمات المفهرسة'
    sheet_head(ws, ['تسلسل', 'السورة', 'اسم السورة', 'الآية', 'موضع الكلمة',
                    'الكلمة', 'الداني', 'السجاوندي'],
               [8, 8, 16, 7, 11, 24, 10, 11])
    rows = con.execute('''
        SELECT w.seq, w.sura, w.sura_name, w.ayah, w.wpos, w.text,
               MAX(CASE WHEN r.source='dani' THEN r.raw END) d,
               MAX(CASE WHEN r.source='saj'  THEN r.raw END) s
        FROM words w LEFT JOIN rulings r ON r.seq = w.seq
        GROUP BY w.seq ORDER BY w.seq''').fetchall()
    for r in rows:
        ws.append([r['seq'], r['sura'], r['sura_name'], r['ayah'], r['wpos'],
                   r['text'], r['d'], r['s']])

    # --- ورقة الأحكام وحدها ---
    ws2 = wb.create_sheet('الأحكام')
    sheet_head(ws2, ['تسلسل', 'الموضع', 'الوقف بعد', 'المصدر', 'الرمز كما في الملف',
                     'المصطلح في المصدر', 'الفئة العامة', 'مطابقة تقريبية', 'ملاحظة', 'للمراجعة'],
               [8, 18, 22, 26, 16, 20, 13, 14, 26, 28])
    for r in con.execute('''
        SELECT r.seq, w.sura_name, w.sura, w.ayah, w.text, s.name src,
               r.raw, r.term, r.category, r.approx, r.note, r.review
        FROM rulings r JOIN words w ON w.seq = r.seq JOIN sources s ON s.id = r.source
        ORDER BY r.seq'''):
        cat = con.execute('SELECT 1').fetchone()
        ws2.append([r['seq'], f"{r['sura_name']} {r['sura']}:{r['ayah']}", r['text'],
                    r['src'], r['raw'], r['term'], r['category'],
                    'نعم' if r['approx'] else 'لا', r['note'], r['review']])

    # --- ورقة قائمة المراجعة ---
    ws3 = wb.create_sheet('قائمة المراجعة')
    sheet_head(ws3, ['تسلسل', 'الموضع', 'الكلمة', 'المصدر', 'القيمة في الملف', 'سبب المراجعة',
                     'القرار (يملؤه المراجع)'],
               [8, 20, 22, 26, 16, 32, 26])
    for r in con.execute('''
        SELECT r.seq, w.sura_name, w.sura, w.ayah, w.text, s.name src, r.raw, r.review
        FROM rulings r JOIN words w ON w.seq = r.seq JOIN sources s ON s.id = r.source
        WHERE r.review IS NOT NULL ORDER BY r.seq'''):
        ws3.append([r['seq'], f"{r['sura_name']} {r['sura']}:{r['ayah']}", r['text'],
                    r['src'], r['raw'], r['review'], ''])
    for row in ws3.iter_rows(min_row=2, max_col=7):
        row[6].fill = NOTE

    for s in wb.worksheets:
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name=AR, size=11)
                c.alignment = Alignment(horizontal='right', vertical='center')

    wb.save(outfile)
    print(f'حُفظ {outfile}: {len(rows)} كلمة، {ws2.max_row - 1} حكماً، {ws3.max_row - 1} للمراجعة')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
