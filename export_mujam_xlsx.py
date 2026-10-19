#!/usr/bin/env python3
"""ورقة تحقّق: كل مدخل بشاهده من المتن، ليُراجَع ما أُثبت في ملف الوورد."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

d = json.load(open('out/mujam_index.json', encoding='utf-8'))
wb = Workbook()
ws = wb.active
ws.title = 'المصنفات'
ws.sheet_view.rightToLeft = True

HEAD = ['#', 'الباب', 'القرن', 'المصنَّف', 'المؤلف', 'منسوب',
        'الحال المستنتجة', 'كل ما وُجد', 'صفحة المعجم', 'الشاهد من المتن']
ws.append(HEAD)
for i, c in enumerate(HEAD, 1):
    cell = ws.cell(row=1, column=i)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='1E3A5F')
    cell.alignment = Alignment(horizontal='center', vertical='center')

for n, e in enumerate(d['entries'], 1):
    ev = e.get('evidence') or []
    ws.append([
        n,
        (e.get('bab') or '')[:40],
        e.get('century') or '',
        e.get('title') or '',
        e.get('author') or '',
        'نعم' if e.get('attributed') else '',
        e.get('label') or 'لم يُعثر على تصريح',
        ' · '.join(e.get('state') or []),
        e.get('body_page') or '',
        ev[0]['quote'][:500] if ev else '',
    ])

widths = [5, 26, 7, 50, 40, 8, 18, 22, 12, 90]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
ws.freeze_panes = 'A2'

ws2 = wb.create_sheet('المنهج')
ws2.sheet_view.rightToLeft = True
for line in [
    ['ورقة تحقّق لملف «مصنفات الوقف والابتداء عبر القرون»'],
    [''],
    ['المصدر', 'معجم مصنفات الوقف والابتداء — د. محمد توفيق حديد — مركز تفسير — ط١ ١٤٣٧هـ/٢٠١٦م'],
    ['طريقة الاستخراج', 'فهرس الموضوعات قُرئ بالتعرّف الضوئي، فاستُخرجت منه المداخل بعناوينها ومؤلفيها.'],
    ['تحديد الموضع', 'يُبحث عن عنوان المدخل في متن المعجم كله، ويُشترط أن يقارنه اسمُ مؤلفه.'],
    ['الحال', 'أولُ لفظٍ صريح من ألفاظ المؤلف في نطاق المدخل، والشاهدُ مثبتٌ في العمود الأخير.'],
    ['ما لا شاهد له', 'يُترك «لم يُعثر على تصريح»، ولا يُملأ بالظنّ.'],
    ['تنبيه', 'النصّ مستخرجٌ آلياً من صورة، ففيه خطأٌ لا محالة؛ والأصلُ هو الحَكَم.'],
]:
    ws2.append(line)
ws2.column_dimensions['A'].width = 20
ws2.column_dimensions['B'].width = 100
for row in ws2.iter_rows():
    for c in row:
        c.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
ws2['A1'].font = Font(bold=True, size=14, color='1E3A5F')

wb.save('out/شواهد_مصنفات_الوقف.xlsx')
print('تمّ:', len(d['entries']), 'صفاً')
